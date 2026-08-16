#!/usr/bin/env python3
"""Independent release verifier for AIAP v6.5.

Run from 08_VERIFICATION:
    python -B verify_release.py ..

The verifier is read-only. It validates the extracted release tree, including
OOXML integrity, navigation, PDF destinations, workbook cached formula values,
machine-readable records, source fidelity, metadata, and cryptographic
manifests. It does not claim empirical validation of AIAP.
"""
from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import subprocess
import sys
import zipfile
import os
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from xml.etree import ElementTree as ET

import yaml
from docx import Document
from docx.oxml.ns import qn
from jsonschema import Draft202012Validator
from pypdf import PdfReader

VERSION = "6.5"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
X_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS = {"w": W_NS, "r": R_NS, "pr": REL_NS, "x": X_NS}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def rel(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


# Narrow allow-list of runtime/editor artifacts that are not release content.
# Governed publication files remain subject to exact manifest and SHA-256 coverage.
_EPHEMERAL_DIRS = {"__pycache__", ".git"}
_EPHEMERAL_NAMES = {".DS_Store", "Thumbs.db"}
_EPHEMERAL_SUFFIXES = {".pyc", ".pyo", ".tmp", ".bak"}


def is_ephemeral(path: Path, root: Path) -> bool:
    try:
        parts = path.relative_to(root).parts
    except ValueError:
        return False
    if any(part in _EPHEMERAL_DIRS for part in parts[:-1]):
        return True
    name = path.name
    return (
        name in _EPHEMERAL_NAMES
        or name.startswith("~$")
        or path.suffix.casefold() in _EPHEMERAL_SUFFIXES
    )


def governed_files(root: Path):
    return (p for p in root.rglob("*") if p.is_file() and not is_ephemeral(p, root))


def add_check(checks: list[dict[str, Any]], name: str, passed: bool, details: Any = None) -> None:
    checks.append({"name": name, "pass": bool(passed), "details": details})


def flatten_outline(reader: PdfReader, items: list[Any] | None = None) -> list[tuple[str, int]]:
    out: list[tuple[str, int]] = []
    for item in (reader.outline if items is None else items) or []:
        if isinstance(item, list):
            out.extend(flatten_outline(reader, item))
        else:
            try:
                page = reader.get_destination_page_number(item) + 1
            except Exception:
                continue
            out.append((getattr(item, "title", str(item)).strip(), page))
    return out


def normalize_title(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text).strip().casefold()


def resolve_pdf_destination(reader: PdfReader, dest: Any) -> bool:
    if dest is None:
        return False
    try:
        if isinstance(dest, str):
            named = reader.named_destinations.get(dest)
            return named is not None and reader.get_destination_page_number(named) >= 0
        if isinstance(dest, list) and dest:
            page_ref = dest[0]
            try:
                return reader.get_page_number(page_ref) >= 0
            except Exception:
                obj = page_ref.get_object() if hasattr(page_ref, "get_object") else page_ref
                return any(page.get_object() == obj for page in reader.pages)
        if hasattr(dest, "get"):
            return reader.get_destination_page_number(dest) >= 0
    except Exception:
        return False
    return False


def inspect_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path), strict=True)
    internal = external = broken = 0
    destinations: set[str] = set()
    for page in reader.pages:
        for ref_obj in page.get("/Annots") or []:
            try:
                annot = ref_obj.get_object()
            except Exception:
                continue
            if str(annot.get("/Subtype")) != "/Link":
                continue
            action = annot.get("/A")
            if action:
                try:
                    action = action.get_object()
                except Exception:
                    pass
                kind = str(action.get("/S"))
                if kind == "/URI":
                    external += 1
                    continue
                if kind == "/GoTo":
                    internal += 1
                    dest = action.get("/D")
                    destinations.add(repr(dest))
                    if not resolve_pdf_destination(reader, dest):
                        broken += 1
                    continue
            dest = annot.get("/Dest")
            if dest is not None:
                internal += 1
                destinations.add(repr(dest))
                if not resolve_pdf_destination(reader, dest):
                    broken += 1
    outline = flatten_outline(reader)
    metadata = {str(k): str(v) for k, v in (reader.metadata or {}).items()}
    return {
        "pages": len(reader.pages),
        "encrypted": bool(reader.is_encrypted),
        "internal_links": internal,
        "distinct_internal_destinations": len(destinations),
        "external_links": external,
        "broken_internal_destinations": broken,
        "outline_entries": len(outline),
        "outline": outline,
        "metadata": metadata,
    }


def inspect_docx(path: Path) -> dict[str, Any]:
    with path.open("rb") as fh:
        magic = fh.read(4).hex(" ")
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        names = set(zf.namelist())
        xml_names = [n for n in names if n.endswith(".xml")]
        xml_text = "\n".join(zf.read(n).decode("utf-8", "ignore") for n in xml_names)
        tracked = len(re.findall(r"<w:(?:ins|del|moveFrom|moveTo)(?:\s|>)", xml_text))
        comments_parts = sorted(n for n in names if n.startswith("word/comments"))
        document_xml = ET.fromstring(zf.read("word/document.xml"))
        bookmark_names = {
            node.get(f"{{{W_NS}}}name")
            for node in document_xml.findall(".//w:bookmarkStart", NS)
            if node.get(f"{{{W_NS}}}name")
        }
        anchors = [
            node.get(f"{{{W_NS}}}anchor")
            for node in document_xml.findall(".//w:hyperlink", NS)
            if node.get(f"{{{W_NS}}}anchor")
        ]
        unresolved = sorted({a for a in anchors if a not in bookmark_names})
        media = sorted(n for n in names if n.startswith("word/media/"))
        wp_ns = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        docpr = document_xml.findall(f".//{{{wp_ns}}}docPr")
        alt_count = sum(1 for n in docpr if (n.get("descr") or n.get("title")))
        external_image_rels: list[str] = []
        external_links: list[str] = []
        for rel_name in [n for n in names if n.startswith("word/_rels/") and n.endswith(".rels")]:
            rroot = ET.fromstring(zf.read(rel_name))
            for r in rroot:
                rtype = r.get("Type") or ""
                target = r.get("Target", "")
                if rtype.endswith("/image") and r.get("TargetMode") == "External":
                    external_image_rels.append(target)
                if rtype.endswith("/hyperlink") and r.get("TargetMode") == "External":
                    external_links.append(target)
    doc = Document(str(path))
    toc: list[tuple[str, int]] = []
    for paragraph in doc.paragraphs:
        style = (paragraph.style.name if paragraph.style else "").casefold()
        if style.startswith("toc") and "\t" in paragraph.text:
            title, page = paragraph.text.rsplit("\t", 1)
            try:
                toc.append((title.strip(), int(page.strip())))
            except ValueError:
                pass
    full_text = "\n".join(p.text for p in doc.paragraphs)
    return {
        "magic": magic,
        "zip_crc_failure": bad,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "media_parts": len(media),
        "drawing_alt_text_count": alt_count,
        "comments_parts": comments_parts,
        "tracked_change_elements": tracked,
        "bookmarks": len(bookmark_names),
        "internal_hyperlinks": len(anchors),
        "unresolved_internal_anchors": unresolved,
        "external_image_relationships": external_image_rels,
        "external_hyperlinks": external_links,
        "materialized_toc": toc,
        "text": full_text,
    }


def toc_sync(docx_info: dict[str, Any], pdf_info: dict[str, Any]) -> dict[str, Any]:
    outline = {normalize_title(title): page for title, page in pdf_info.get("outline", [])}
    mismatches: list[dict[str, Any]] = []
    missing: list[str] = []
    for title, old_page in docx_info.get("materialized_toc", []):
        actual = outline.get(normalize_title(title))
        if actual is None:
            missing.append(title)
        elif actual != old_page:
            mismatches.append({"title": title, "toc_page": old_page, "actual_page": actual})
    return {"toc_entries": len(docx_info.get("materialized_toc", [])), "missing_outline_targets": missing, "page_mismatches": mismatches}


def workbook_parts(zf: zipfile.ZipFile) -> dict[str, str]:
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    rel_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rel_map = {node.get("Id"): node.get("Target", "").lstrip("/") for node in rel_root}
    out: dict[str, str] = {}
    sheets_node = wb.find(f"{{{X_NS}}}sheets")
    for sheet in ([] if sheets_node is None else list(sheets_node)):
        rid = sheet.get(f"{{{R_NS}}}id")
        if rid in rel_map:
            out[sheet.get("name", "")] = rel_map[rid]
    return out


def style_formats(zf: zipfile.ZipFile) -> list[str]:
    root = ET.fromstring(zf.read("xl/styles.xml"))
    custom = {
        int(node.get("numFmtId")): node.get("formatCode", "")
        for node in root.findall(f".//{{{X_NS}}}numFmt")
    }
    builtins = {0: "General", 1: "0", 2: "0.00", 9: "0%", 10: "0.00%"}
    xfs = root.find(f"{{{X_NS}}}cellXfs")
    if xfs is None:
        return []
    return [custom.get(int(xf.get("numFmtId", "0")), builtins.get(int(xf.get("numFmtId", "0")), f"builtin:{xf.get('numFmtId')}")) for xf in xfs]


def inspect_xlsx(path: Path) -> dict[str, Any]:
    with path.open("rb") as fh:
        magic = fh.read(4).hex(" ")
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        names = set(zf.namelist())
        formulas = 0
        formulas_with_cache = 0
        error_cells: list[str] = []
        formula_cache_missing: list[str] = []
        old_version_hits: list[str] = []
        for name in sorted(n for n in names if n.startswith("xl/") and n.endswith(".xml")):
            data = zf.read(name)
            text = data.decode("utf-8", "ignore")
            if any(token in text for token in ("6.0" + ".1", "6.0" + ".2")):
                old_version_hits.append(name)
            if name.startswith("xl/worksheets/sheet"):
                root = ET.fromstring(data)
                for cell in root.findall(f".//{{{X_NS}}}c"):
                    ref = cell.get("r", "?")
                    f = cell.find(f"{{{X_NS}}}f")
                    v = cell.find(f"{{{X_NS}}}v")
                    if f is not None:
                        formulas += 1
                        if v is not None and v.text not in (None, ""):
                            formulas_with_cache += 1
                        else:
                            formula_cache_missing.append(f"{name}:{ref}")
                    if cell.get("t") == "e":
                        error_cells.append(f"{name}:{ref}={v.text if v is not None else ''}")
        calc_attrs: dict[str, str] = {}
        wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
        calc = wb_root.find(f"{{{X_NS}}}calcPr")
        if calc is not None:
            calc_attrs = dict(calc.attrib)
        return {
            "magic": magic,
            "zip_crc_failure": bad,
            "formula_count": formulas,
            "formula_cached_values": formulas_with_cache,
            "missing_formula_cached_values": formula_cache_missing,
            "cached_error_cells": error_cells,
            "calc_properties": calc_attrs,
            "old_version_xml_parts": old_version_hits,
        }


def workload_format_checks(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as zf:
        sheets = workbook_parts(zf)
        formats = style_formats(zf)
        inputs = ET.fromstring(zf.read(sheets["Inputs"]))
        cells = {c.get("r"): c for c in inputs.findall(f".//{{{X_NS}}}c")}
        expected = {
            "B8": "0%", "C8": "0%", "B9": "0%", "C9": "0%", "B10": "0%", "C10": "0%",
            "B13": "0%", "C13": "0%", "B15": "0%", "C15": "0%", "B18": "0%", "C18": "0%",
            "B22": "0%", "C22": "0%", "B16": "0", "C16": "0", "B20": "0.00", "C20": "0.00",
            "B21": "0.00", "C21": "0.00", "B23": "0.00", "C23": "0.00",
        }
        actual: dict[str, str] = {}
        values: dict[str, str] = {}
        for ref, wanted in expected.items():
            cell = cells.get(ref)
            if cell is None:
                actual[ref] = "MISSING"
                continue
            style_id = int(cell.get("s", "0"))
            actual[ref] = formats[style_id] if style_id < len(formats) else f"style:{style_id}"
            values[ref] = cell.findtext(f"{{{X_NS}}}v", default="")
        mismatches = {ref: {"expected": expected[ref], "actual": actual.get(ref)} for ref in expected if actual.get(ref) != expected[ref]}
        value_checks = {
            "B16_minutes": values.get("B16") == "10",
            "C16_minutes": values.get("C16") == "10",
            "B20_hours": values.get("B20") == "3",
            "C20_hours": values.get("C20") == "3",
        }
        return {"format_mismatches": mismatches, "value_checks": value_checks, "actual_formats": actual}


def feasibility_formula_check(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as zf:
        formulas = "\n".join(zf.read(n).decode("utf-8", "ignore") for n in zf.namelist() if n.startswith("xl/worksheets/") and n.endswith(".xml"))
    return {
        "requires_all_13_nonblank": "COUNTA(D5:D17)&lt;13" in formulas,
        "affirmative_13_pass": 'COUNTIF(D5:D17,"PASS")=13' in formulas,
        "residual_pass_pattern_absent": 'PASS_WITH_CONTROLS","PASS"' not in formulas,
    }


def docx_visible_text(path: Path) -> str:
    """Return paragraph and table-cell text from a DOCX for content assertions."""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def xlsx_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    values: list[str] = []
    for si in root.findall(f"{{{X_NS}}}si"):
        values.append("".join(t.text or "" for t in si.findall(f".//{{{X_NS}}}t")))
    return values


def xlsx_sheet_cells(path: Path, sheet_name: str) -> tuple[dict[str, dict[str, Any]], ET.Element]:
    """Read visible values and formulas from a named worksheet without Excel."""
    with zipfile.ZipFile(path) as zf:
        sheets = workbook_parts(zf)
        if sheet_name not in sheets:
            raise KeyError(f"worksheet absent: {sheet_name}")
        shared = xlsx_shared_strings(zf)
        sheet_root = ET.fromstring(zf.read(sheets[sheet_name]))
        cells: dict[str, dict[str, Any]] = {}
        for cell in sheet_root.findall(f".//{{{X_NS}}}c"):
            ref = cell.get("r", "")
            ctype = cell.get("t")
            formula = cell.findtext(f"{{{X_NS}}}f")
            if ctype == "inlineStr":
                value = "".join(t.text or "" for t in cell.findall(f".//{{{X_NS}}}t"))
            else:
                raw = cell.findtext(f"{{{X_NS}}}v")
                if ctype == "s" and raw not in (None, ""):
                    try:
                        value = shared[int(raw)]
                    except Exception:
                        value = raw
                elif ctype == "b":
                    value = raw == "1"
                else:
                    value = raw
            cells[ref] = {"value": value, "formula": formula, "type": ctype}
        return cells, sheet_root


def operational_hardening_checks(root: Path) -> dict[str, Any]:
    """Verify the bounded final corrections arising from the Claude/Genspark audit."""
    findings: dict[str, Any] = {}
    failures: list[str] = []

    # Pilot Data Dictionary: the released instrument must be able to run RQ1-RQ3.
    dictionary = root / "04_PILOT_AND_VALIDATION" / "AIAP_Pilot_Data_Dictionary_v6.5.xlsx"
    dict_cells, _ = xlsx_sheet_cells(dictionary, "Data Dictionary")
    fields: dict[str, dict[str, str]] = {}
    for row in range(4, 200):
        field = str(dict_cells.get(f"B{row}", {}).get("value") or "").strip()
        if field:
            fields[field] = {
                "domain": str(dict_cells.get(f"A{row}", {}).get("value") or ""),
                "required": str(dict_cells.get(f"F{row}", {}).get("value") or ""),
                "rq": str(dict_cells.get(f"G{row}", {}).get("value") or ""),
                "notes": str(dict_cells.get(f"H{row}", {}).get("value") or ""),
            }
    required_fields = {
        "participant_id", "programme_id", "module_id", "eal_status",
        "accommodation_status", "assessment_anxiety_score", "tool_access",
        "assessment_id", "claim_id", "lane", "modality", "occasion_id",
        "product_score", "authentication_event_id", "authentication_total_score",
        "authentication_decision", "authentication_duration_minutes",
        "unseen_or_perturbed", "submission_authentication_interval_minutes",
        "prompt_id", "form_id", "prompt_source", "exposure_status",
        "independent_criterion_score", "independent_criterion_decision",
        "criterion_timing", "criterion_assessor_blinded", "alternative_route_used",
        "student_preparation_minutes", "student_wait_setup_minutes",
        "student_event_minutes", "staff_role_minutes", "capacity_incident",
        "privacy_incident",
    }
    missing_fields = sorted(required_fields - set(fields))
    criterion_conditional = fields.get("independent_criterion_score", {}).get("required") == "Conditional"
    findings["pilot_data_dictionary"] = {
        "field_count": len(fields),
        "missing_required_analysis_fields": missing_fields,
        "independent_criterion_is_conditional": criterion_conditional,
    }
    if missing_fields:
        failures.append(f"Pilot Data Dictionary missing fields: {missing_fields}")
    if not criterion_conditional:
        failures.append("independent_criterion_score must be Conditional")

    # Workload comparator: invigilation is concurrent, not cohort-multiplied.
    workload = root / "05_WORKBOOKS" / "AIAP_Workload_Calculator_v6.5.xlsx"
    load_cells, _ = xlsx_sheet_cells(workload, "Local Comparator")
    formula_d16 = (load_cells.get("D16", {}).get("formula") or "").replace(" ", "")
    def as_float(ref: str) -> float | None:
        value = load_cells.get(ref, {}).get("value")
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
    comparator_ok = (
        formula_d16 == "B7*B6/60*B8+B9"
        and as_float("B7") == 1.0
        and as_float("D16") == 8.0
        and as_float("D19") == 113.0
        and as_float("D20") is not None
        and abs(as_float("D20") - (113.0 / 300.0)) < 1e-9
    )
    findings["workload_local_comparator"] = {
        "D16_formula": formula_d16,
        "number_of_sittings": as_float("B7"),
        "invigilation_hours": as_float("D16"),
        "recurring_total": as_float("D19"),
        "hours_per_student": as_float("D20"),
        "pass": comparator_ok,
    }
    if not comparator_ok:
        failures.append("Workload Local Comparator correction failed")

    # Prompt construct blueprint must default to an unassessed, fail-closed state.
    prompt = root / "04_PILOT_AND_VALIDATION" / "AIAP_Prompt_Bank_and_Exposure_Log_v6.5.xlsx"
    prompt_cells, prompt_root = xlsx_sheet_cells(prompt, "Construct Blueprint")
    prompt_defaults = [prompt_cells.get(f"C{row}", {}).get("value") for row in range(4, 54)]
    validation_text = " ".join(
        node.findtext(f"{{{X_NS}}}formula1", default="")
        for node in prompt_root.findall(f".//{{{X_NS}}}dataValidation")
    )
    prompt_ok = all(value == "Not assessed" for value in prompt_defaults) and "Yes,No,Not assessed" in validation_text
    findings["prompt_construct_blueprint"] = {
        "not_assessed_defaults": sum(value == "Not assessed" for value in prompt_defaults),
        "expected_rows": len(prompt_defaults),
        "validation": validation_text,
        "pass": prompt_ok,
    }
    if not prompt_ok:
        failures.append("Prompt Construct Blueprint is not fail-closed")

    # Feasibility workbook serialization maps every display and machine token.
    feasibility = root / "05_WORKBOOKS" / "AIAP_Assurance_Feasibility_Model_v6.5.xlsx"
    feas_cells, _ = xlsx_sheet_cells(feasibility, "Guidance")
    mapping_text = " ".join(str(feas_cells.get(ref, {}).get("value") or "") for ref in ("A10", "B10", "C10"))
    mapping_tokens = ["PASS", "PASS_WITH_CONTROLS", "FAIL", "NOT READY", "NOT_READY", "Not assessed", "NOT_ASSESSED"]
    mapping_ok = all(token in mapping_text for token in mapping_tokens)
    findings["status_serialization"] = {"text": mapping_text, "pass": mapping_ok}
    if not mapping_ok:
        failures.append("Feasibility status serialization mapping incomplete")

    # Programme workbook must make the exemplar's non-normative status visible.
    programme = root / "05_WORKBOOKS" / "AIAP_Programme_Assurance_Map_v6.5.xlsx"
    programme_cells, _ = xlsx_sheet_cells(programme, "Programme Map")
    programme_note = str(programme_cells.get("A2", {}).get("value") or "")
    programme_ok = "not a recommended default" in programme_note
    findings["programme_exemplar_boundary"] = {"text": programme_note, "pass": programme_ok}
    if not programme_ok:
        failures.append("Programme Assurance Map exemplar boundary missing")

    # DOCX/source synchronization and exact normative corrections.
    core_docx = root / "01_CORE_PAPER" / "AIAP_Core_Working_Paper_v6.5.docx"
    core_text = docx_visible_text(core_docx)
    core_checks = {
        "evidence_date_14_august": "verified as of 14 August 2026" in core_text,
        "stale_13_august_absent": "verified as of 13 August 2026" not in core_text,
        "competence_authentication_includes_verify": "explain, defend, verify, adapt, or transfer" in core_text,
        "table_16a_phase_1": "Phase 1 — RQ1 and RQ2" in core_text,
        "table_16a_phase_2": "Phase 2 — RQ2 / architecture-comparison" in core_text,
    }
    findings["core_docx_corrections"] = core_checks
    if not all(core_checks.values()):
        failures.append("Core DOCX final synchronization/correction check failed")

    standard = root / "02_NORMATIVE_STANDARD" / "AIAP_Normative_Assurance_Standard_v6.5.docx"
    standard_text = docx_visible_text(standard)
    gate_labels = [
        "G01 — Claim map approved",
        "G02 — Prompt security governed",
        "G03 — Parallel-form equivalence supported",
        "G04 — Submission-authentication interval justified",
        "G05 — Assessor qualification and calibration",
        "G06 — Reviewer and panel capacity",
        "G07 — Independence and conflicts controlled",
        "G08 — Accessibility-equivalent route operational",
        "G09 — Appeal and second-review route operational",
        "G10 — Data governance approved",
        "G11 — Capacity-degradation fallback declared",
        "G12 — Threat model acceptable",
        "G13 — No unresolved high-severity exception",
    ]
    individual_rule = (
        "every student SHALL complete the declared minimum authentication" in standard_text
        and "sampling SHALL NOT substitute for individual certification evidence" in standard_text
    )
    standard_ok = all(label in standard_text for label in gate_labels) and individual_rule
    findings["normative_standard"] = {
        "gate_labels_present": {label: label in standard_text for label in gate_labels},
        "individual_certification_shall": individual_rule,
        "pass": standard_ok,
    }
    if not standard_ok:
        failures.append("Normative Standard G01-G13 or individual-certification rule incomplete")

    readiness = root / "03_IMPLEMENTATION" / "AIAP_Pilot_Readiness_Pack_v6.5.docx"
    readiness_text = docx_visible_text(readiness)
    readiness_checks = {
        "pass_only_all_13": "PASS is returned only when all thirteen conditions are PASS" in readiness_text,
        "controlled_only_no_fail_unassessed": (
            "PASS_WITH_CONTROLS is returned only when at least one condition is PASS_WITH_CONTROLS"
            in readiness_text
            and "none is FAIL or NOT_ASSESSED" in readiness_text
        ),
        "other_not_ready": "Any other state is NOT_READY" in readiness_text,
    }
    findings["pilot_readiness_formula"] = readiness_checks
    if not all(readiness_checks.values()):
        failures.append("Pilot Readiness Pack overall decision formula remains ambiguous")

    # Canonical Markdown byte-level encoding and accepted content.
    source = root / "01_CORE_PAPER" / "AIAP_Core_Working_Paper_v6.5_SOURCE.md"
    source_bytes = source.read_bytes()
    try:
        source_text = source_bytes.decode("utf-8")
        utf8 = True
    except UnicodeDecodeError:
        source_text = ""
        utf8 = False
    unicode_tokens = ["Kılınç", "Gašević", "Raković", "Foltýnek", "Šigut", "–", "×", "§", "©"]
    unicode_ok = utf8 and all(token in source_text for token in unicode_tokens)
    findings["markdown_encoding"] = {
        "utf8_decode": utf8,
        "non_ascii_bytes": sum(byte >= 128 for byte in source_bytes),
        "tokens": {token: token in source_text for token in unicode_tokens},
        "pass": unicode_ok,
    }
    if not unicode_ok:
        failures.append("Canonical Markdown Unicode/UTF-8 verification failed")

    findings["failures"] = failures
    findings["pass"] = not failures
    return findings


def inspect_markdown(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    anchors = set(re.findall(r'<a\s+id="([^"]+)"\s*></a>', text))
    links = re.findall(r"\]\(#([^)]+)\)", text)
    unresolved = sorted({link for link in links if link not in anchors})
    images = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text)
    missing_images = [target for target in images if not (path.parent / target.split("#", 1)[0]).exists()]
    exact_tokens = [
        "Kılınç", "Gašević", "Raković", "Foltýnek", "Šigut", "Levels 2–5", "R0–R6",
        "10–15%", "50–100-word", "300 × 5 minutes", "Standard §11", "©", "“oralness”",
    ]
    refs = text.split("# References", 1)[1] if "# References" in text else ""
    return {
        "anchors": len(anchors),
        "internal_links": len(links),
        "unresolved_internal_links": unresolved,
        "image_references": images,
        "missing_images": missing_images,
        "exact_tokens_present": {token: token in text for token in exact_tokens},
        "appendix_d3_anchor": "app-d3" in anchors,
        "appendix_d3_contents": "[Appendix D3." in text,
        "reference_blank_gap_absent": "\n\n\n" not in refs,
        "plain_correspondence": "Correspondence: [james.mg@buv.edu.vn]" in text,
        "phase_sequence": "Its empirical programme proceeds in two explicit phases" in text,
        "claim_status_rationale": "Claim status is first-class in AIAP" in text,
        "misconduct_alignment": "In Lane 2A, outsourcing is misconduct only when the brief has explicitly made production authorship or a production constraint part of the claim" in text,
        "luo_reference": "10.1080/02602938.2024.2309963" in text,
        "dawson_reference": "10.1080/02602938.2017.1336746" in text,
        "live_aias_faq_path": "https://aiassessmentscale.com/faqs/" in text,
        "stale_aias_faq_path": "aiassessmentscale.com/260-2" in text,
        "thirteen_gate_phrase": "thirteen fail-closed conditions" in text,
        "stale_seven_gate_phrase": "requires all seven conditions" in text,
        "evidence_date_14_august": "verified as of 14 August 2026" in text,
        "stale_evidence_date_13_absent": "verified as of 13 August 2026" not in text,
        "competence_authentication_includes_verify": "explain, defend, verify, adapt, or transfer" in text,
        "phase_1_table_16a": "Phase 1 — RQ1 and RQ2" in text,
        "phase_2_table_16a": "Phase 2 — RQ2 / architecture-comparison" in text,
    }


def inspect_cff(path: Path) -> dict[str, Any]:
    obj = yaml.safe_load(path.read_text(encoding="utf-8"))
    preferred = obj.get("preferred-citation") or {}
    errors: list[str] = []
    if obj.get("cff-version") != "1.2.0": errors.append("cff-version must be 1.2.0")
    if obj.get("version") != VERSION: errors.append("root version must be 6.5")
    if obj.get("date-released") != "2026-08-14": errors.append("release date mismatch")
    if obj.get("type") not in {"software", "dataset"}: errors.append("root type must be software or dataset")
    for field in ("message", "title", "authors", "url"):
        if not obj.get(field): errors.append(f"missing root field: {field}")
    if preferred.get("type") != "report": errors.append("preferred citation must be report")
    if "v6.5" not in str(preferred.get("version")): errors.append("preferred citation version mismatch")
    return {"errors": errors, "root_type": obj.get("type"), "preferred_type": preferred.get("type")}


def html_link_check(path: Path, root: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    hrefs = re.findall(r'href="([^"]+)"', text)
    missing: list[str] = []
    for href in hrefs:
        parsed = urlparse(href)
        if parsed.scheme or href.startswith("#"):
            continue
        target = (path.parent / parsed.path).resolve()
        try:
            target.relative_to(root)
        except ValueError:
            missing.append(href)
            continue
        if not target.exists():
            missing.append(href)
    return {"hrefs": hrefs, "missing_relative_targets": missing, "canonical_present": 'https://ripplelogic.org/aiap/' in text}


def main() -> int:
    root = Path(sys.argv[1] if len(sys.argv) > 1 else "..").resolve()
    checks: list[dict[str, Any]] = []
    errors: list[str] = []
    details: dict[str, Any] = {}

    required = {
        ".gitignore", "README.md", "START_HERE.md", "VERSION_MANIFEST.yaml", "CHANGELOG.md",
        "RELEASE_NOTES_v6.5.md", "LICENSE.md", "RIGHTS_AND_LICENSING.md", "CITATION.cff",
        "CONTRIBUTING.md", "SECURITY.md", "requirements.txt",
        "01_CORE_PAPER/AIAP_Core_Working_Paper_v6.5.docx",
        "01_CORE_PAPER/AIAP_Core_Working_Paper_v6.5.pdf",
        "01_CORE_PAPER/AIAP_Core_Working_Paper_v6.5_SOURCE.md",
        "01_CORE_PAPER/figures/AIAP_Figure_1_At_A_Glance_v6.5.png",
        "01_CORE_PAPER/figures/AIAP_Figure_2_Evidence_Plan_Architecture_v6.5.png",
        "01_CORE_PAPER/figures/AIAP_Figure_3_Independence_Matrix_v6.5.png",
        "01_CORE_PAPER/figures/AIAP_Figure_4_Forbidden_Inferential_Transfers_v6.5.png",
        "01_CORE_PAPER/figures/AIAP_Figure_5_Programme_Assurance_Spine_v6.5.png",
        "02_NORMATIVE_STANDARD/AIAP_Normative_Assurance_Standard_v6.5.docx",
        "02_NORMATIVE_STANDARD/AIAP_Normative_Assurance_Standard_v6.5.pdf",
        "02_NORMATIVE_STANDARD/AIAP_v6.5_CORE_TO_STANDARD_CROSSWALK.md",
        "03_IMPLEMENTATION/AIAP_Implementation_Handbook_v6.5.docx",
        "03_IMPLEMENTATION/AIAP_Implementation_Handbook_v6.5.pdf",
        "03_IMPLEMENTATION/AIAP_Pilot_Readiness_Pack_v6.5.docx",
        "03_IMPLEMENTATION/AIAP_Pilot_Readiness_Pack_v6.5.pdf",
        "04_PILOT_AND_VALIDATION/AIAP_Pilot_and_Validation_Protocol_v6.5.docx",
        "04_PILOT_AND_VALIDATION/AIAP_Pilot_and_Validation_Protocol_v6.5.pdf",
        "04_PILOT_AND_VALIDATION/AIAP_Pilot_Data_Dictionary_v6.5.xlsx",
        "04_PILOT_AND_VALIDATION/AIAP_Prompt_Bank_and_Exposure_Log_v6.5.xlsx",
        "05_WORKBOOKS/AIAP_Workload_Calculator_v6.5.xlsx",
        "05_WORKBOOKS/AIAP_Assurance_Feasibility_Model_v6.5.xlsx",
        "05_WORKBOOKS/AIAP_Programme_Assurance_Map_v6.5.xlsx",
        "06_MACHINE_READABLE/AIAP_Assessment_Assurance_Record_v6.5.schema.json",
        "06_MACHINE_READABLE/AIAP_Assessment_Assurance_Record_v6.5.example.json",
        "06_MACHINE_READABLE/AIAP_Assessment_Assurance_Record_v6.5.example_all_lanes.json",
        "06_MACHINE_READABLE/validate_aiap_record.py",
        "06_MACHINE_READABLE/test_validate_aiap_record.py",
        "07_PUBLICATION/AIAP_RippleLogic_Page_v6.5.html",
        "07_PUBLICATION/AIAP_RippleLogic_Metadata_v6.5.json",
        "07_PUBLICATION/AIAP_v6.5_SSRN_Metadata.md",
        "07_PUBLICATION/AIAP_v6.5_SSRN_Abstract.txt",
        "08_VERIFICATION/AIAP_v6.5_Feedback_Adjudication_and_Change_Log.md",
        "08_VERIFICATION/BUILD_AND_TEST_REPORT.json", "08_VERIFICATION/VERIFICATION_STATUS.md",
        "08_VERIFICATION/FILE_MANIFEST.csv", "08_VERIFICATION/SHA256SUMS.txt", "08_VERIFICATION/verify_release.py",
    }
    missing = sorted(item for item in required if not (root / item).exists())
    add_check(checks, "required publication and repository inventory", not missing, {"missing": missing, "required_count": len(required)})
    errors.extend(f"missing required file: {item}" for item in missing)

    # SHA-256 manifest covers every release file except itself.
    hash_path = root / "08_VERIFICATION" / "SHA256SUMS.txt"
    hash_errors: list[str] = []
    hash_entries: dict[str, str] = {}
    if hash_path.exists():
        for line in hash_path.read_text(encoding="utf-8").splitlines():
            if not line.strip(): continue
            if "  " not in line:
                hash_errors.append(f"malformed line: {line}")
                continue
            expected, item = line.split("  ", 1)
            hash_entries[item] = expected
            p = root / item
            if not p.exists(): hash_errors.append(f"missing: {item}")
            elif sha256(p) != expected: hash_errors.append(f"hash mismatch: {item}")
        expected_items = {rel(p, root) for p in governed_files(root) if p != hash_path}
        hash_errors.extend(f"unhashed: {item}" for item in sorted(expected_items - set(hash_entries)))
        hash_errors.extend(f"extra: {item}" for item in sorted(set(hash_entries) - expected_items))
    else:
        hash_errors.append("SHA256SUMS.txt absent")
    add_check(checks, "SHA-256 manifest exact coverage and verification", not hash_errors, {"entries": len(hash_entries), "errors": hash_errors})
    errors.extend(hash_errors)

    # File manifest covers every file except the two manifests to avoid recursion.
    file_manifest_path = root / "08_VERIFICATION" / "FILE_MANIFEST.csv"
    manifest_errors: list[str] = []
    rows: list[dict[str, str]] = []
    if file_manifest_path.exists():
        with file_manifest_path.open(newline="", encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        listed = {row.get("path", "") for row in rows}
        expected_items = {
            rel(p, root) for p in governed_files(root)
            if p not in {file_manifest_path, hash_path}
        }
        manifest_errors.extend(f"unlisted: {item}" for item in sorted(expected_items - listed))
        manifest_errors.extend(f"extra: {item}" for item in sorted(listed - expected_items))
        for row in rows:
            p = root / row["path"]
            if not p.exists(): continue
            if int(row["bytes"]) != p.stat().st_size: manifest_errors.append(f"size mismatch: {row['path']}")
            if row["sha256"] != sha256(p): manifest_errors.append(f"hash mismatch: {row['path']}")
    else:
        manifest_errors.append("FILE_MANIFEST.csv absent")
    add_check(checks, "file manifest exact coverage, size, and hash verification", not manifest_errors, {"rows": len(rows), "errors": manifest_errors})
    errors.extend(manifest_errors)

    # DOCX and PDF integrity plus TOC synchronization.
    docx_results: dict[str, Any] = {}
    pdf_results: dict[str, Any] = {}
    toc_results: dict[str, Any] = {}
    expected_pages = {
        "01_CORE_PAPER/AIAP_Core_Working_Paper_v6.5.pdf": 87,
        "02_NORMATIVE_STANDARD/AIAP_Normative_Assurance_Standard_v6.5.pdf": 8,
        "03_IMPLEMENTATION/AIAP_Data_Protection_Accessibility_Evidence_Governance_v6.5.pdf": 4,
        "03_IMPLEMENTATION/AIAP_Implementation_Handbook_v6.5.pdf": 6,
        "03_IMPLEMENTATION/AIAP_Pilot_Readiness_Pack_v6.5.pdf": 8,
        "03_IMPLEMENTATION/AIAP_Role_Coding_Manual_v6.5.pdf": 5,
        "03_IMPLEMENTATION/AIAP_Student_and_Instructor_Toolkit_v6.5.pdf": 5,
        "04_PILOT_AND_VALIDATION/AIAP_Pilot_and_Validation_Protocol_v6.5.pdf": 6,
    }
    doc_pdf_errors: list[str] = []
    for docx in sorted(p for p in root.rglob("*.docx") if not is_ephemeral(p, root)):
        rpath = rel(docx, root)
        try:
            info = inspect_docx(docx)
            docx_results[rpath] = {k: v for k, v in info.items() if k != "text"}
            if info["magic"] != "50 4b 03 04" or info["zip_crc_failure"] is not None:
                doc_pdf_errors.append(f"invalid OOXML: {rpath}")
            if info["comments_parts"] or info["tracked_change_elements"]:
                doc_pdf_errors.append(f"review residue: {rpath}")
            if info["unresolved_internal_anchors"] or info["external_image_relationships"]:
                doc_pdf_errors.append(f"broken/linked content: {rpath}")
            if any("](" in target or "aiassessmentscale.com/260-2" in target for target in info["external_hyperlinks"]):
                doc_pdf_errors.append(f"malformed or stale external hyperlink relationship: {rpath}")
            if docx.name == "AIAP_Core_Working_Paper_v6.5.docx":
                if info["media_parts"] != 5 or info["drawing_alt_text_count"] != 5:
                    doc_pdf_errors.append("core must contain five embedded figures with alt text")
                if "Correspondence: james.mg@buv.edu.vn" not in info["text"] or "Correspondence: mailto:" in info["text"]:
                    doc_pdf_errors.append("core correspondence display is not normalized")
                if "mailto:james.mg@buv.edu.vn" not in info["external_hyperlinks"]:
                    doc_pdf_errors.append("core correspondence mailto target missing")
        except Exception as exc:
            doc_pdf_errors.append(f"DOCX parse failure {rpath}: {exc}")
    for pdf in sorted(p for p in root.rglob("*.pdf") if not is_ephemeral(p, root)):
        rpath = rel(pdf, root)
        try:
            info = inspect_pdf(pdf)
            pdf_results[rpath] = info
            if info["encrypted"] or info["broken_internal_destinations"]:
                doc_pdf_errors.append(f"PDF destination/encryption failure: {rpath}")
            if info["outline_entries"] == 0:
                doc_pdf_errors.append(f"PDF outline absent: {rpath}")
            if rpath in expected_pages and info["pages"] != expected_pages[rpath]:
                doc_pdf_errors.append(f"PDF page count drift: {rpath} ({info['pages']} != {expected_pages[rpath]})")
            if info["metadata"].get("/Author") != "James McGaughran":
                doc_pdf_errors.append(f"PDF author metadata mismatch: {rpath}")
            if "6.5" not in (info["metadata"].get("/Subject", "") + info["metadata"].get("/Title", "")):
                doc_pdf_errors.append(f"PDF version metadata mismatch: {rpath}")
        except Exception as exc:
            doc_pdf_errors.append(f"PDF parse failure {rpath}: {exc}")
    for docx_rel, d_info in docx_results.items():
        pdf_rel = str(Path(docx_rel).with_suffix(".pdf")).replace("\\", "/")
        if pdf_rel not in pdf_results:
            doc_pdf_errors.append(f"matching PDF absent for {docx_rel}")
            continue
        sync = toc_sync(d_info, pdf_results[pdf_rel])
        toc_results[docx_rel] = sync
        if sync["missing_outline_targets"] or sync["page_mismatches"]:
            doc_pdf_errors.append(f"TOC/page synchronization failure: {docx_rel}")
        toc_count = sync["toc_entries"]
        if toc_count and pdf_results[pdf_rel]["internal_links"] < toc_count:
            doc_pdf_errors.append(f"insufficient PDF TOC links: {pdf_rel}")
    core_pdf = pdf_results.get("01_CORE_PAPER/AIAP_Core_Working_Paper_v6.5.pdf", {})
    if core_pdf and (core_pdf.get("internal_links", 0) < 175 or core_pdf.get("distinct_internal_destinations", 0) < 70 or core_pdf.get("external_links", 0) < 80):
        doc_pdf_errors.append("core PDF link structure below expected floor")
    add_check(checks, "DOCX/PDF integrity, metadata, live navigation, embedded figures, and materialized TOC synchronization", not doc_pdf_errors, {"errors": doc_pdf_errors, "docx": docx_results, "pdf": pdf_results, "toc": toc_results})
    errors.extend(doc_pdf_errors)
    details["docx"] = docx_results
    details["pdf"] = pdf_results
    details["toc"] = toc_results

    # XLSX structure, caches, formulas, and corrected formats.
    xlsx_results: dict[str, Any] = {}
    xlsx_errors: list[str] = []
    for xlsx in sorted(p for p in root.rglob("*.xlsx") if not is_ephemeral(p, root)):
        rpath = rel(xlsx, root)
        try:
            info = inspect_xlsx(xlsx)
            xlsx_results[rpath] = info
            if info["magic"] != "50 4b 03 04" or info["zip_crc_failure"] is not None:
                xlsx_errors.append(f"invalid OOXML: {rpath}")
            if info["cached_error_cells"] or info["missing_formula_cached_values"]:
                xlsx_errors.append(f"formula cache/error defect: {rpath}")
            if info["old_version_xml_parts"]:
                xlsx_errors.append(f"stale version in workbook XML: {rpath}")
        except Exception as exc:
            xlsx_errors.append(f"XLSX parse failure {rpath}: {exc}")
    workload = root / "05_WORKBOOKS" / "AIAP_Workload_Calculator_v6.5.xlsx"
    workload_info = workload_format_checks(workload)
    xlsx_results[rel(workload, root)]["workload_format_checks"] = workload_info
    if workload_info["format_mismatches"] or not all(workload_info["value_checks"].values()):
        xlsx_errors.append("Workload Calculator unit/number formats failed")
    feasibility = root / "05_WORKBOOKS" / "AIAP_Assurance_Feasibility_Model_v6.5.xlsx"
    feasibility_info = feasibility_formula_check(feasibility)
    xlsx_results[rel(feasibility, root)]["affirmative_pass_formula"] = feasibility_info
    if not all(feasibility_info.values()):
        xlsx_errors.append("Assurance Feasibility Model is not affirmative/fail-closed")
    add_check(checks, "XLSX OOXML integrity, formula caches, fail-closed gate, and corrected unit formats", not xlsx_errors, {"errors": xlsx_errors, "workbooks": xlsx_results})
    errors.extend(xlsx_errors)
    details["xlsx"] = xlsx_results

    # Final bounded operational-hardening corrections from the Claude/Genspark audit.
    hardening = operational_hardening_checks(root)
    add_check(
        checks,
        "final operational hardening: pilot dataset, workload comparator, normative gates, fail-closed defaults, and cross-artifact synchronization",
        hardening["pass"],
        hardening,
    )
    if not hardening["pass"]:
        errors.extend(f"operational hardening: {item}" for item in hardening["failures"])
    details["operational_hardening"] = hardening

    # Machine-readable schema, examples, semantic validator, and adversarial tests.
    machine = root / "06_MACHINE_READABLE"
    machine_errors: list[str] = []
    machine_details: dict[str, Any] = {}
    try:
        schema = json.loads((machine / "AIAP_Assessment_Assurance_Record_v6.5.schema.json").read_text(encoding="utf-8"))
        Draft202012Validator.check_schema(schema)
        if schema.get("$id") != "https://ripplelogic.org/aiap/v6.5/assessment-assurance-record.schema.json":
            machine_errors.append("schema $id mismatch")
        if schema.get("properties", {}).get("protocol_version", {}).get("const") != VERSION:
            machine_errors.append("schema version const mismatch")
        validator = Draft202012Validator(schema)
        validation: dict[str, list[str]] = {}
        for name in ("AIAP_Assessment_Assurance_Record_v6.5.example.json", "AIAP_Assessment_Assurance_Record_v6.5.example_all_lanes.json"):
            obj = json.loads((machine / name).read_text(encoding="utf-8"))
            validation[name] = [err.message for err in validator.iter_errors(obj)]
        machine_details["schema_examples"] = validation
        if any(validation.values()): machine_errors.append("schema example validation failed")
    except Exception as exc:
        machine_errors.append(f"schema failure: {exc}")
    child_env = dict(os.environ)
    child_env["PYTHONDONTWRITEBYTECODE"] = "1"
    cp = subprocess.run([sys.executable, str(machine / "validate_aiap_record.py"), str(machine / "AIAP_Assessment_Assurance_Record_v6.5.example.json"), str(machine / "AIAP_Assessment_Assurance_Record_v6.5.example_all_lanes.json")], capture_output=True, text=True, env=child_env)
    machine_details["semantic_validator"] = {"returncode": cp.returncode, "stdout": cp.stdout, "stderr": cp.stderr}
    if cp.returncode: machine_errors.append("semantic validator failed")
    cp2 = subprocess.run([sys.executable, "-m", "unittest", "-v", "test_validate_aiap_record.py"], cwd=machine, capture_output=True, text=True, env=child_env)
    machine_details["unit_tests"] = {"returncode": cp2.returncode, "stdout": cp2.stdout, "stderr": cp2.stderr}
    if cp2.returncode: machine_errors.append("adversarial unit tests failed")
    add_check(checks, "JSON Schema, two examples, semantic validator, and nine adversarial tests", not machine_errors, {"errors": machine_errors, **machine_details})
    errors.extend(machine_errors)
    details["machine"] = machine_details

    # Canonical source fidelity and the accepted scholarly corrections.
    source_info = inspect_markdown(root / "01_CORE_PAPER" / "AIAP_Core_Working_Paper_v6.5_SOURCE.md")
    source_ok = (
        not source_info["unresolved_internal_links"] and not source_info["missing_images"]
        and all(source_info["exact_tokens_present"].values())
        and source_info["appendix_d3_anchor"] and source_info["appendix_d3_contents"]
        and source_info["reference_blank_gap_absent"] and source_info["plain_correspondence"]
        and source_info["phase_sequence"] and source_info["claim_status_rationale"]
        and source_info["misconduct_alignment"] and source_info["luo_reference"] and source_info["dawson_reference"]
        and source_info["live_aias_faq_path"] and not source_info["stale_aias_faq_path"]
        and source_info["thirteen_gate_phrase"] and not source_info["stale_seven_gate_phrase"]
        and source_info["evidence_date_14_august"] and source_info["stale_evidence_date_13_absent"]
        and source_info["competence_authentication_includes_verify"]
        and source_info["phase_1_table_16a"] and source_info["phase_2_table_16a"]
    )
    add_check(checks, "canonical UTF-8 source, anchors, figures, reference integrity, and accepted v6.5 corrections", source_ok, source_info)
    if not source_ok: errors.append("canonical source fidelity/correction check failed")
    details["source"] = source_info

    # Repository and publication metadata.
    metadata_errors: list[str] = []
    manifest = yaml.safe_load((root / "VERSION_MANIFEST.yaml").read_text(encoding="utf-8"))
    if str(manifest.get("release", {}).get("version")) != VERSION: metadata_errors.append("VERSION_MANIFEST version mismatch")
    if str(manifest.get("release", {}).get("date")) != "2026-08-14": metadata_errors.append("VERSION_MANIFEST date mismatch")
    cff = inspect_cff(root / "CITATION.cff")
    metadata_errors.extend(f"CFF: {item}" for item in cff["errors"])
    license_text = (root / "LICENSE.md").read_text(encoding="utf-8")
    rights_text = (root / "RIGHTS_AND_LICENSING.md").read_text(encoding="utf-8")
    if "All rights reserved" not in license_text or "All rights reserved" not in rights_text:
        metadata_errors.append("rights notice missing")
    html_info = html_link_check(root / "07_PUBLICATION" / "AIAP_RippleLogic_Page_v6.5.html", root)
    if html_info["missing_relative_targets"] or not html_info["canonical_present"]:
        metadata_errors.append("RippleLogic HTML link/canonical failure")
    web_meta = json.loads((root / "07_PUBLICATION" / "AIAP_RippleLogic_Metadata_v6.5.json").read_text(encoding="utf-8"))
    expected_release_url = "https://github.com/MathGov/ripple-logic/releases/tag/AIAP-v6.5"
    if (
        web_meta.get("version") != VERSION
        or web_meta.get("canonical_url") != "https://ripplelogic.org/aiap/"
        or web_meta.get("github_release_url") != expected_release_url
        or web_meta.get("release_zip") != "AIAP_v6.5_COMPLETE_READY_FINAL.zip"
    ):
        metadata_errors.append("RippleLogic metadata mismatch")
    html_text = (root / "07_PUBLICATION" / "AIAP_RippleLogic_Page_v6.5.html").read_text(encoding="utf-8")
    for required_url in (
        expected_release_url,
        "https://github.com/MathGov/ripple-logic/releases/download/AIAP-v6.5/AIAP_Core_Working_Paper_v6.5.pdf",
        "https://github.com/MathGov/ripple-logic/releases/download/AIAP-v6.5/AIAP_Normative_Assurance_Standard_v6.5.pdf",
        "https://github.com/MathGov/ripple-logic/releases/download/AIAP-v6.5/AIAP_v6.5_COMPLETE_READY_FINAL.zip",
    ):
        if required_url not in html_text:
            metadata_errors.append(f"RippleLogic HTML missing canonical release URL: {required_url}")
    add_check(checks, "GitHub-root metadata, CFF, rights, and RippleLogic.org publication package", not metadata_errors, {"errors": metadata_errors, "cff": cff, "html": html_info, "web_metadata": web_meta})
    errors.extend(metadata_errors)
    details["publication_metadata"] = {"cff": cff, "html": html_info, "web_metadata": web_meta}

    # Stale version and placeholder scan. Historical files may name prior releases.
    stale: list[str] = []
    placeholders: list[str] = []
    text_suffixes = {".md", ".txt", ".json", ".yaml", ".yml", ".cff", ".py", ".csv", ".html"}
    historical_allow = {"CHANGELOG.md", "08_VERIFICATION/AIAP_v6.5_Feedback_Adjudication_and_Change_Log.md"}
    placeholder_re = re.compile(r"\b(TODO|TBD|FIXME|PLACEHOLDER)\b|\[INSERT", re.I)
    old_tokens = ("6.0" + ".1", "6.0" + ".2")
    for path in sorted(governed_files(root)):
        rpath = rel(path, root)
        if path.suffix.casefold() in text_suffixes:
            if rpath in historical_allow or path.name in {"verify_release.py", "BUILD_AND_TEST_REPORT.json", "SHA256SUMS.txt", "FILE_MANIFEST.csv"}:
                continue
            try: text = path.read_text(encoding="utf-8")
            except Exception: continue
            if any(token in text for token in old_tokens): stale.append(rpath)
            if placeholder_re.search(text): placeholders.append(rpath)
        elif path.suffix.casefold() in {".docx", ".xlsx"}:
            with zipfile.ZipFile(path) as zf:
                for name in zf.namelist():
                    if not name.endswith((".xml", ".rels")): continue
                    text = zf.read(name).decode("utf-8", "ignore")
                    if any(token in text for token in old_tokens):
                        stale.append(f"{rpath}!{name}")
                        break
        elif path.suffix.casefold() == ".pdf":
            text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
            if any(token in text for token in old_tokens): stale.append(rpath)
    add_check(checks, "stale-version and placeholder scan", not stale and not placeholders, {"stale": stale, "placeholders": placeholders})
    errors.extend(f"stale prior-version token: {item}" for item in stale)
    errors.extend(f"placeholder token: {item}" for item in placeholders)

    counts = Counter((p.suffix.casefold() or "[no extension]") for p in governed_files(root))
    summary = {
        "release": f"AIAP v{VERSION}",
        "root": str(root),
        "pass": not errors,
        "checks_passed": sum(1 for check in checks if check["pass"]),
        "checks_total": len(checks),
        "file_count": sum(1 for _ in governed_files(root)),
        "file_types": dict(sorted(counts.items())),
        "errors": errors,
        "checks": checks,
        "details": details,
    }
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
